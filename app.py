from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Any
from urllib.parse import quote
import httpx
import json
from io import BytesIO
from openpyxl import load_workbook
import os
import sys

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# PyInstaller 빌드 시 sys._MEIPASS, 개발 시 파일 위치 사용
if getattr(sys, "frozen", False):
    BASE_DIR = sys._MEIPASS
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(BASE_DIR, "부록설계_양식.xlsx")

SYSTEM_PROMPT = """당신은 이벤트·공공사업 분야 최상위 제안서 전략가입니다.
업로드된 제안요청서(RFP)/과업지시서 원문을 심층 분석하고,
웹 검색으로 발주처의 최근 사업 이력·평가 성향·유사 행사 수주 사례를 조사한 뒤,
아래 JSON 스키마에 정확히 맞춰 결과만 출력하세요.
설명, 코드블록 표시(백틱), 그 외 텍스트는 절대 포함하지 마세요. 순수 JSON 객체만 출력합니다.

[분석 5원칙]
1. 사실 기반: 문서에 명시된 내용을 최우선으로 인용하고, 웹 검색으로 보완합니다. 문서에 없는 내용을 창작하거나 추측하는 것은 금지합니다.
2. 숨은 니즈 발굴: 명시적 요구사항 너머 발주처가 진짜 원하는 것(정치적 이해관계, 성과 KPI, 리스크 회피 심리)을 파악해 전략에 반영합니다.
3. 평가 기준 우선: 평가 항목과 배점을 문서에서 반드시 추출합니다. 배점이 높은 항목일수록 핵심과업·차별화포인트·슬라이드 분량에서 비중을 높게 설계합니다.
4. 사업 유형 판단: marketing_pr / event / it_system / public / consulting 중 가장 유사한 유형으로 분류하고, 해당 유형의 수주 공식을 전략에 적용합니다.
5. Win Theme 3원칙: 경쟁사 대비 수주 가능성이 가장 높은 핵심 전략 메시지(Win Theme) 3개를 정의하고, 이를 제안작성방향과 목차설계 전반에 일관되게 반영합니다.

[발주처정보 작성 기준]
- 기관특성: 기관 규모·성격, 최근 행정·정치 이슈, 선호 파트너사 스타일, 전년도 사업 평가 성향까지 포함합니다.
- 최근사업이력: "연도 · 사업명 · 수행사 · 예산 · 특이사항" 형식으로 최소 3건 이상 웹 검색해 작성합니다.
- 출처는 실제 URL 또는 기관 공고 페이지명을 기재합니다.

[과거유사사례 작성 기준]
- 동일 발주처 또는 유사 규모·성격의 행사·사업 수주 사례를 3건 이상 조사합니다.
- 수주사의 차별화 전략, 수주 금액, 경쟁 포인트를 구체적으로 기술합니다.
- 수주 실패 또는 재입찰 사례가 있다면 이유와 함께 기재합니다.

[제안작성방향 작성 기준]
- 핵심전략: Win Theme 3개를 하나의 수주 전략 메시지로 통합한 문장으로 작성합니다.
- 차별화포인트: 경쟁사 대비 구체적 강점을 수치·수행 이력·보유 자원과 함께 작성합니다.
- 유의사항: 배점 비중이 높은 필수 대응 항목, 발주처의 민감 이슈, 자격 미달·감점 리스크 요인을 구체적으로 작성합니다.

[목차설계 작성 기준]
- 대분류는 반드시 아래 4개 중 하나를 사용합니다. 임의로 변경하거나 다른 값을 사용하는 것은 금지합니다.
  "제안 개요"     → Phase 0-2: 오프닝·훅, Win Theme 요약, 사업이해·시장인사이트 (최대 4개 항목)
  "제안사 소개"   → Phase 6: 수행 역량, 투입 인력 및 조직, 유사 수행 실적 (최대 3개 항목)
  "사업수행계획"  → Phase 3-5·7: 전략개념, 실행계획, 품질·리스크 관리, 기대효과·ROI (최대 18개 항목)
  "사업 관리"     → 계약·정산·보고·성과품 관리 등 행정·관리 항목 (최대 5개 항목)
- 슬라이드제목(소분류)은 액션 중심으로 작성합니다.
- 슬라이드핵심내용은 C-E-I 구조로 작성합니다: 주장(Claim) → 근거(Evidence) → 임팩트(Impact)
- 각 항목은 아래 필드로 구성합니다.
  · 소분류: 해당 섹션의 소제목
  · 주요기획: 해당 섹션에서 다룰 핵심 기획 방향을 2~3문장으로 서술
  · 포함내용: 해당 슬라이드/섹션에 반드시 포함해야 할 세부 요소·내용 항목을 열거
  · 필수조건: RFP 또는 평가 기준에서 해당 섹션과 직접 관련된 필수 반영 조건
  · 헤드메세지: 슬라이드 최상단에 배치할 한 문장 Action Title. 숫자·성과 중심으로 작성
- Impact-8 프레임워크 순서를 따릅니다:
  Phase 0 훅·오프닝 → Phase 1 요약·Win Theme → Phase 2 시장인사이트 → Phase 3 전략개념
  → Phase 4 실행계획(전체의 30~45% 분량) → Phase 5 관리·품질보증 → Phase 6 수행역량 → Phase 7 투자·기대효과·ROI
- 전체 항목 수는 15~25개 이내로 실제 제출 가능한 수준으로 구성합니다.

스키마:
{
  "사업개요": { "사업명": "", "발주처": "", "사업목적": "", "사업기간": "", "예산": "", "핵심과업": [""] },
  "발주처정보": { "기관특성": "", "최근사업이력": [""], "참고출처": [""] },
  "과거유사사례": [ { "사례명": "", "주관사": "", "특이사항": "", "참고출처": "" } ],
  "제안작성방향": { "핵심전략": "", "차별화포인트": [""], "유의사항": [""] },
  "목차설계": [ { "대분류": "", "소분류": "", "주요기획": "", "포함내용": "", "필수조건": "", "헤드메세지": "" } ]
}"""


class AnalyzeRequest(BaseModel):
    text: str
    api_key: str


class ExportRequest(BaseModel):
    analysis: dict[str, Any]


@app.post("/analyze")
async def analyze(req: AnalyzeRequest):
    truncated = req.text[:60000]
    async with httpx.AsyncClient(timeout=120.0) as client:
        resp = await client.post(
            f"https://generativelanguage.googleapis.com/v1beta/models/gemini-3.6-flash:generateContent?key={req.api_key}",
            json={
                "systemInstruction": {"parts": [{"text": SYSTEM_PROMPT}]},
                "contents": [{"role": "user", "parts": [{"text": f"다음은 제안요청서 원문입니다. 분석해서 JSON으로만 응답하세요.\n\n---\n{truncated}\n---"}]}],
                "generationConfig": {"maxOutputTokens": 16000},
            },
        )

    data = resp.json()

    if "error" in data:
        raise HTTPException(400, data["error"].get("message", "Gemini API 오류"))
    if data.get("promptFeedback", {}).get("blockReason"):
        raise HTTPException(400, f"안전 필터 차단: {data['promptFeedback']['blockReason']}")

    candidates = data.get("candidates") or []
    if not candidates:
        raise HTTPException(500, "후보 응답 없음")

    candidate = candidates[0]
    finish_reason = candidate.get("finishReason")
    if finish_reason == "SAFETY":
        raise HTTPException(400, "안전 정책으로 응답이 차단됐습니다.")
    if finish_reason == "MAX_TOKENS":
        raise HTTPException(400, "응답이 토큰 제한으로 잘렸습니다.")

    parts = candidate.get("content", {}).get("parts", [])
    raw = "".join(p["text"] for p in parts if "text" in p).strip()
    cleaned = raw.replace("```json", "").replace("```", "").strip()

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        raise HTTPException(500, f"JSON 파싱 실패: {cleaned[:300]}")


@app.post("/export")
async def export_excel(req: ExportRequest):
    if not os.path.exists(TEMPLATE_PATH):
        raise HTTPException(500, f"템플릿 파일을 찾을 수 없습니다: {TEMPLATE_PATH}")

    analysis = req.analysis
    o = analysis.get("사업개요", {})
    toc = analysis.get("목차설계", [])

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    ws["A1"].value = f"{o.get('사업명', '')} — 제안서 부록 설계"
    ws["B2"].value = o.get("사업명", "")
    ws["B3"].value = o.get("예산", "")
    ws["B4"].value = o.get("사업기간", "")

    # JS의 0-based 행 인덱스를 openpyxl 1-based로 변환 (+1)
    SECTIONS = [
        {"name": "제안 개요",    "start": 7,  "end": 10},
        {"name": "제안사 소개",  "start": 11, "end": 13},
        {"name": "사업수행계획", "start": 14, "end": 31},
        {"name": "사업 관리",   "start": 32, "end": 36},
    ]

    grouped: dict[str, list] = {}
    for item in toc:
        grouped.setdefault(item.get("대분류", ""), []).append(item)

    for sec in SECTIONS:
        for i, item in enumerate(grouped.get(sec["name"], [])):
            row = sec["start"] + i
            if row > sec["end"]:
                break
            ws.cell(row=row, column=2).value = item.get("소분류", "")
            ws.cell(row=row, column=3).value = item.get("주요기획", "")
            ws.cell(row=row, column=4).value = item.get("포함내용", "")
            ws.cell(row=row, column=5).value = item.get("필수조건", "")
            ws.cell(row=row, column=6).value = item.get("헤드메세지", "")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{o.get('사업명', 'RFP')}_부록설계.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@app.get("/")
async def serve_index():
    return FileResponse(os.path.join(BASE_DIR, "index.html"))
